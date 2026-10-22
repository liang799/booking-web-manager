"""Build the editable session test report from reviewed cases and run evidence.

Requires python-docx and Pillow. Run from any directory:
  python build_report.py --output session-testing-report.docx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
WIDTH = 10.48


def setup(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
    sec.top_margin = sec.bottom_margin = Inches(0.52)
    sec.left_margin = sec.right_margin = Inches(0.60)
    sec.footer_distance = Inches(0.24)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 27), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(8)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Booking Web Manager  |  Session testing  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Black box and white box testing for session use cases"
    doc.core_properties.subject = "UC2-01 UC2-02 UC2-03a UC2-03b UC2-03c"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "testing, equivalence partitioning, boundary values, control flow"


def p(doc, text, size=None, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return para


def heading(doc, title, new_page=True):
    if new_page:
        doc.add_page_break()
    doc.add_heading(title, level=1)


def table(doc, headers, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    properties = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    properties.append(borders)
    for row_index, values in enumerate([headers, *rows]):
        row = t.rows[0] if row_index == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        prevent_split = OxmlElement("w:cantSplit")
        trpr.append(prevent_split)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            trpr.append(repeat)
        for i, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "75")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tcpr.append(margins)
            fill = OxmlElement("w:shd")
            fill.set(qn("w:fill"), "E7E7E7" if row_index == 0 else "FFFFFF")
            tcpr.append(fill)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.04
            para.paragraph_format.keep_with_next = row_index == 0
            if i == 0 or headers[i] in ["Result", "Actual output", "Passed", "Tests"]:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(value))
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = row_index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def figure(doc, path, max_height=3.9):
    image = Image.open(path)
    w, h = image.size
    physical_w = min(10.1, max_height * w / h)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline = para.add_run().add_picture(str(path), width=Inches(physical_w))
    inline._inline.docPr.set("descr", path.stem.replace("-", " ") + " control flow graph")


def result_text(value):
    if value["outcome"] == "created":
        return f"Created {value['shareCents']}c; saved {value['persistedSessions']}"
    return f"{value['error']}; saved {value['persistedSessions']}"


def changed_input(row, baseline):
    parts = []
    for group in ["booking", "config"]:
        current = row["input"][group]
        for key, value in current.items():
            if baseline["input"][group].get(key) != value:
                parts.append(f"{key}={value}")
    for key in ["actor", "now"]:
        if row["input"][key] != baseline["input"][key]:
            parts.append(f"{key}={row['input'][key]}")
    return "; ".join(parts) if parts else "Baseline; optional values omitted"


def create_report(output):
    cases = json.loads((HERE / "manual-cases.json").read_text())
    graphs = json.loads((HERE / "figures/graphs.json").read_text())
    observations = json.loads((HERE / "evidence/testing-black-box-observations.json").read_text())["observations"]
    doc = Document()
    setup(doc)
    doc.add_paragraph("Black box and white box testing\nfor session use cases", style="Title")
    p(doc, "Booking Web Manager   |   SC2006 Group 3   |   20 October 2026", 12)
    p(doc, "This report covers Neoh's session discovery, creation and management use cases. It provides user-facing test scenarios, code-based control-flow diagrams and recorded automated results that can be used in the project's testing section.")
    production = json.loads((HERE / "evidence/production/execution-summary.json").read_text())
    p(doc, production["opening_summary"], bold=True)
    table(doc, ["Use case", "Function", "Black box focus", "White box subject"], [
        ["UC2-01", "Discover sessions", "Public visibility, search, filters and pagination", "SQL filter construction"],
        ["UC2-02", "Create session", "Account eligibility, capacity, dates and price", "Account eligibility helper"],
        ["UC2-03a", "Change session visibility", "Owner, capacity, lifecycle and explicit target", "Visibility input parser"],
        ["UC2-03b", "Remove participant", "Preview, historical refund, replay and state", "Removal confirmation callback"],
        ["UC2-03c", "Cancel session", "Refunds, retained history, preview and replay", "Cancellation preparation helper"],
    ], [1.0, 1.6, 3.8, 4.08], 10)
    p(doc, "Black-box design selects inputs from valid and invalid groups and around boundaries. The expected result is judged through a public interface. The creation appendix applies those techniques at the typed application interface; its test fixture also inspects saved state and financial instructions.", 10)
    p(doc, "White-box design selects paths after inspecting implementation decisions. Each graph states its scope and counting convention. A helper's complexity is not the complexity of the entire use case, and source-derived paths are not instrumented coverage measurements.", 10)
    p(doc, "Implementation baseline: 135064c0244fc684c9773ff947d8bf44e60f085e. The inspected session source and the existing focused tests are unchanged in PR base 7aacdc9. Shared domain helpers are included because they implement these use cases; ownership does not imply sole authorship.", 9)
    p(doc, "Reference format: the supplied senior SRS.pdf, section 6.9, PDF pages 77-106. Its TripSmart functionality and pass results are not requirements or evidence for Booking Web Manager.", 9)

    for group in cases:
        heading(doc, f"Black box testing {group['uc']}")
        doc.add_heading(group["title"], level=2)
        p(doc, f"Scenario: {group['scenario']}  Priority: High.", 9.5)
        p(doc, "Prerequisites: " + group["prerequisites"], 9.5)
        rows = [[r["id"], r["description"], r["action"], r["inputs"], r["expected"],
                 r.get("actual_summary", r.get("actual", "Not run")),
                 r.get("result", "Pending")] for r in group["cases"]]
        table(doc, ["ID", "Description", "Action", "Inputs", "Expected output", "Actual output", "Result"],
              rows, [0.48, 1.1, 1.5, 1.65, 3.05, 1.87, 0.83], 8.6)
        p(doc, "Pass covers all listed subcases at the stated interface. Partial means only some were observed. Blocked means a required account, fixture or controlled environment is unavailable. The production execution section gives the observations and limits; local unit results do not turn a production case into a pass.", 9)
        p(doc, "Source: " + group["source"], 8.5)

    heading(doc, "Production execution results")
    p(doc, production["run_context"], 10)
    for paragraph in production["summary_paragraphs"]:
        p(doc, paragraph, 10)
    table(doc, ["Check", "Observed output", "Conclusion"], production["checks"],
          [2.2, 5.2, 3.08], 9)
    p(doc, production["limitations"], 9)
    p(doc, "Evidence: evidence/production/ contains sanitized HTTP results, browser observations and the read-only reproduction script. No account passwords, tokens or private user records are included.", 9)

    heading(doc, "Authenticated browser checks")
    p(doc, production["authenticated_context"], 10)
    table(doc, ["Check", "Observed output", "Conclusion"], production["authenticated_checks"],
          [2.2, 5.2, 3.08], 9)
    p(doc, production["authenticated_limitations"], 9)
    p(doc, "Evidence: evidence/production/authenticated-browser-observations.json and its per-check accessibility snapshots. Account and payout-guard screenshots are included with the evidence.", 9)

    heading(doc, "Production API checks with the test account")
    p(doc, production["api_context"], 10)
    table(doc, ["Check", "Observed output", "Conclusion"], production["api_checks"],
          [2.2, 5.2, 3.08], 9)
    p(doc, production["api_limitations"], 9)
    p(doc, "Evidence: evidence/production/authenticated-api-observations.json retains sanitized request inputs, response status/error codes and timestamps. Authentication used the normal password login; credentials and tokens are excluded.", 9)

    graph_items = graphs["graphs"] if isinstance(graphs, dict) else graphs
    by_image = {g["pngFile"]: g for g in graph_items}
    # The figure manifest records all node and edge details; the narrative gives
    # each graph's boundary and deliberately avoids inferred coverage percentages.
    white = json.loads((HERE / "white-box-cases.json").read_text())
    for item in white:
        graph = by_image[item["image"]]
        assert len(item["paths"]) == graph["counts"]["cyclomaticComplexity"]
        for case, path in zip(item["paths"], graph["paths"], strict=True):
            assert case[1] == " ".join(path["nodes"])
        heading(doc, f"White box testing {item['uc']}")
        doc.add_heading(item["title"], level=2)
        p(doc, "Source: " + item["source"], 9)
        p(doc, item["scope"], 9.5)
        figure(doc, HERE / "figures" / item["image"], max_height=3.85 if item["separate_paths"] else 2.05)
        p(doc, item["complexity"], 10, bold=True)
        if item["separate_paths"]:
            heading(doc, f"Basis paths for {item['uc']}")
            p(doc, item["title"], 11, bold=True)
        table(doc, ["Path", "Basis set of nodes", "Input and expected output", "Execution evidence"], item["paths"],
              [0.55, 3.45, 3.55, 2.93], 9)
        for note in item["notes"]:
            p(doc, note, 9.3)

    for start in range(0, len(observations), 12):
        heading(doc, "Executed creation boundary tests")
        p(doc, f"UC2-02   |   Cases {start + 1}-{min(start + 12, len(observations))} of 46   |   Typed application boundary", 11, bold=True)
        if start == 0:
            p(doc, "Baseline: ready active booker with email and completed payout setup; Jurong East Sports Hall, West, Badminton; cost 1001c; 3 slots; minimum headcount 2; start 1 Nov 2026 18:00 SGT; end 20:00 SGT; clock 30 Oct 2026 18:00 SGT. Chosen price, reliability and visibility are omitted. Dates shown in changed inputs use UTC with Z.", 9)
        else:
            p(doc, "Use the baseline on the first appendix page. Each case varies the listed input. Successful cases also check generated identities and stored settings; every case checks zero creation ledger instructions and zero payout requests.", 9)
        rows = []
        for row in observations[start:start+12]:
            assert row["expected"] == row["actual"]
            rows.append([row["id"], changed_input(row, observations[0]), row["technique"],
                         result_text(row["expected"]), result_text(row["actual"]), "Pass"])
        table(doc, ["ID", "Changed input", "Technique", "Expected output", "Actual output", "Result"],
              rows, [0.5, 3.15, 1.75, 2.25, 2.25, 0.58], 8.5)
        p(doc, "Evidence: evidence/testing-black-box-observations.json; tests/lab4/create-sessions-black-box.test.ts. These are class-boundary results using transaction doubles, not browser or database results.", 8.5)

    heading(doc, "Execution evidence and reproduction")
    p(doc, "The local evidence totals 228 passing checks across 13 files: the original 223-test snapshot at 15:52–15:54 SGT and five isolated discovery basis tests at 16:51 SGT on 20 October 2026. Execution used Node 22.17.0 and Vitest 4.1.11, with no failures or skips in these selections. These outputs remain separate from the production HTTP and browser observations above.", 10)
    table(doc, ["Executed suite", "Tests", "What the result establishes"], [
        ["Discovery application", "4", "Criteria and server-time forwarding; returned projection and error propagation"],
        ["Creation application", "25", "Creation rules, defaults, persistence calls and failure behavior"],
        ["Session management application", "24", "Hosted-session reads; account, ownership, lifecycle and capacity guards"],
        ["Removal application", "30", "Preview validity, refund instructions, replay and rollback in a fake transaction"],
        ["Cancellation application", "19", "Refund/history behavior, replay and fake-transaction failures"],
        ["Creation EP and BVA", "46", "Recorded expected and actual outputs across selected input partitions"],
        ["Cancellation basis cases", "5", "Concrete outcomes for the five predicted helper paths"],
        ["Discovery query and SQL reader", "37", "32 parser checks and 5 SQL-construction/mocked-row checks"],
        ["Isolated discovery basis inputs", "5", "One optional filter per test; exact SQL predicates and bound values"],
        ["Visibility routes actions and OpenAPI", "33", "20 route, 11 action and 2 contract checks"],
        ["Total local focused evidence", "228", "13 distinct files; no real database or external payment execution"],
    ], [3.35, 0.65, 6.48], 9)
    p(doc, "Evidence files retain three timestamped Vitest JSON reports and console logs. The accompanying README provides exact reproduction commands and source links. Tests inspect mocks and in-memory state; PostgreSQL atomicity, concurrent transactions, provider payments and the three-second discovery update target still require their service-backed checks.", 9)
    p(doc, "No branch-coverage percentage, MC/DC result or whole-system pass claim is made. All six predicted discovery basis inputs now have local SQL-construction evidence, including five new isolated filters. The paths remain source-derived predictions, not instrumented traces. Cancellation's caller loop and all helper internals excluded from a graph need their own tests.", 9)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=HERE / "session-testing-report.docx")
    create_report(parser.parse_args().output)
