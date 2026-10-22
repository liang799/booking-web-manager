"""Build the editable account and group test report from reviewed cases and run evidence.

Requires python-docx and Pillow. Run from any directory:
  python build_report.py --output account-testing-report.docx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
BASE_REVISION = "7aacdc97e2ab7d72724429b13eb7f17ee398010b"


def setup(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
    sec.top_margin = sec.bottom_margin = Inches(0.52)
    sec.left_margin = sec.right_margin = Inches(0.60)
    sec.footer_distance = Inches(0.24)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [("Title", 27), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(8)
        doc.styles[name].paragraph_format.space_after = Pt(7)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Booking Web Manager  |  Account and group testing  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Black box and white box testing for account and group use cases"
    doc.core_properties.subject = "UC1-01 UC1-02 UC1-03 UC1-04 UC1-06"
    doc.core_properties.author = ""
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    doc.core_properties.keywords = "testing, equivalence partitioning, boundary values, basis path"


def p(doc, text, size=None, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return para


def heading(doc, title, new_page=True):
    if new_page:
        doc.add_page_break()
    doc.add_heading(title, level=1)


def table(doc, headers, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    # Schema order: tblBorders comes before tblLayout and tblLook.
    anchor = t._tbl.tblPr.find(qn("w:tblLayout"))
    if anchor is None:
        anchor = t._tbl.tblPr.find(qn("w:tblLook"))
    anchor.addprevious(borders)
    for row_index, values in enumerate([headers, *rows]):
        row = t.rows[0] if row_index == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for i, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            tcpr = cell._tc.get_or_add_tcPr()
            # Schema order: tcW, shd, tcMar, vAlign.
            fill = OxmlElement("w:shd")
            fill.set(qn("w:val"), "clear")
            fill.set(qn("w:fill"), "E7E7E7" if row_index == 0 else "FFFFFF")
            tcpr.append(fill)
            margins = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), "75")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tcpr.append(margins)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.04
            para.paragraph_format.keep_with_next = row_index == 0
            if i == 0 or headers[i] in ["Result", "Tests"]:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(value))
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = row_index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def figure(doc, path, max_height, description=None):
    w, h = Image.open(path).size
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline = para.add_run().add_picture(str(path), width=Inches(min(10.1, max_height * w / h)))
    inline._inline.docPr.set("descr", description or path.stem.replace("-", " ") + " control flow graph")


# ---------- Rendering of recorded observations ----------

OBLIGATIONS = {
    "availableBalance": ("available", "c"), "heldBalance": ("held", "c"),
    "activeCommitments": ("commitments", ""), "unsettledOwnedSessions": ("hosted sessions", ""),
    "pendingPayouts": ("pending payouts", ""), "activeOwnedGroups": ("owned groups", ""),
}


def delete_input(row):
    parts = [f"{label}={row['input'][key]}{unit}" for key, (label, unit) in OBLIGATIONS.items() if key in row["input"]]
    if row["input"].get("walletId") is None:
        parts.append("no wallet")
    return "; ".join(parts) if parts else "Baseline: nothing outstanding"


def delete_output(value):
    steps = "check only" if value["steps"] == ["check"] else ", ".join(value["steps"])
    return f"{value['outcome']}; {steps}; profile {value['profileStatus']}"


NOMINAL_PROFILE = {"displayName": "Marcus Lim", "preferredSports": ["Badminton"], "preferredRegions": ["West"]}


def short(value):
    if isinstance(value, list):
        return "[" + ", ".join(value) + "]" if value else "[]"
    text = str(value)
    if text.strip() == "":
        return f'"{text}" ({len(text)} spaces)' if text else '"" (empty)'
    if len(text) > 20:
        return f"{len(text)} chars" + (" incl. spaces" if text != text.strip() else "")
    return f'"{text}"'


def profile_input(row):
    parts = [f"{key}={short(row['input'][key])}" for key in NOMINAL_PROFILE if row["input"][key] != NOMINAL_PROFILE[key]]
    if row["input"]["account"] != "ACTIVE":
        parts.append(f"account={row['input']['account']}")
    return "; ".join(parts) if parts else "Baseline profile, active account"


def profile_output(value):
    if value["outcome"] != "SAVED":
        return f"{value['outcome']}; nothing saved"
    saved = value["saved"]
    changes = [f"{key}={short(saved[key])}" for key in NOMINAL_PROFILE if saved[key] != NOMINAL_PROFILE[key]]
    return "Saved once" + (": " + "; ".join(changes) if changes else ": baseline values")


def executed_tables(doc, title, uc, observations, render_input, render_output, baseline, evidence):
    for start in range(0, len(observations), 12):
        heading(doc, title)
        p(doc, f"{uc}   |   Cases {start + 1}-{min(start + 12, len(observations))} of {len(observations)}   |   Typed application boundary", 11, bold=True)
        p(doc, baseline if start == 0 else "Use the baseline on the first page of this table. Each case changes only the listed input.", 9)
        rows = []
        for row in observations[start:start + 12]:
            assert row["expected"] == row["actual"], row["id"]
            rows.append([row["id"], render_input(row), row["technique"],
                         render_output(row["expected"]), render_output(row["actual"]), "Pass"])
        table(doc, ["ID", "Changed input", "Technique", "Expected output", "Actual output", "Result"],
              rows, [0.5, 2.6, 2.0, 2.4, 2.4, 0.58], 8.5)
        p(doc, "Evidence: " + evidence + " These are use-case results with in-memory ports, not browser or database results.", 8.5)


def create_report(output):
    cases = json.loads((HERE / "manual-cases.json").read_text())
    graphs = {g["id"]: g for g in json.loads((HERE / "figures/graphs.json").read_text())["graphs"]}
    white = json.loads((HERE / "white-box-cases.json").read_text())
    evidence = HERE / "evidence"
    delete_obs = json.loads((evidence / "delete-account-black-box-observations.json").read_text())["observations"]
    profile_obs = json.loads((evidence / "update-profile-black-box-observations.json").read_text())["observations"]
    basis_obs = {o["id"]: o for name in ["account-basis-path-observations.json", "auth-profile-basis-path-observations.json"]
                 for o in json.loads((evidence / name).read_text())["observations"]}
    metadata = json.loads(next(evidence.glob("account-unit-*.metadata.json")).read_text())
    run = json.loads(next(p for p in evidence.glob("account-unit-*.json") if not p.name.endswith(".metadata.json")).read_text())
    totals = metadata["totals"]
    manual_count = sum(len(group["cases"]) for group in cases)
    results = [r["result"] for group in cases for r in group["cases"]]
    production = json.loads((evidence / "production/browser-run-20261020.json").read_text())

    doc = Document()
    setup(doc)
    doc.add_paragraph("Black box and white box testing\nfor account and group use cases", style="Title")
    p(doc, "Booking Web Manager   |   SC2006 Group 3   |   20 October 2026", 12)
    p(doc, "This report covers Joseph's account and group use cases: registration, login, profile settings, account deletion and regular groups. It provides user-facing test scenarios, executed black-box tables, control-flow graphs with basis paths, and recorded automated results.")
    p(doc, f"The focused automated checks passed {totals['passed']} tests across {totals['files']} files, with {totals['todo']} recorded as todo and none failed. "
           f"{len(delete_obs) + len(profile_obs) + len(basis_obs)} of them are the Lab 4 cases in this report, each with recorded expected and actual output. "
           f"The {manual_count} browser scenarios were run on the production site on 20 October 2026: {results.count('Pass')} passed, "
           f"{results.count('Partial')} passed in part and {results.count('Blocked')} were blocked.", bold=True)
    p(doc, "Partial means some subcases matched while others could not be run or did not match the expected output; each mismatch is listed under Defects found. "
           "Blocked means the run could not create the required data: production confirms email at sign-up, and no money or sessions were created.", 10)
    table(doc, ["Use case", "Function", "Black box focus", "White box subject"], [
        ["UC1-01", "Register user", "Password, name and email partitions; duplicate account; invitation return", "registerUser server action (7 paths)"],
        ["UC1-02", "Authenticate user", "Credentials, unconfirmed and deleted accounts, protected pages, redirects", "safeRedirectPath after login (8 paths)"],
        ["UC1-03", "Manage profile", "Executed: name length, sport and region lists, account state (19 cases)", "prepareProfileUpdate policy (10 paths)"],
        ["UC1-04", "Delete account", "Executed: six obligations around zero, invalid facts (21 cases)", "deleteAccount use case (5 paths)"],
        ["UC1-06", "Manage group", "Name limits, links, owner-only actions, archive rule", "Joining through an invitation link (5 paths)"],
    ], [1.0, 1.6, 4.9, 2.98], 10)
    p(doc, "Black-box design selects inputs from valid and invalid groups (equivalence partitioning) and on either side of each limit (boundary value analysis). The expected result is judged only through the class's public interface: the returned outcome, the steps it asked its ports to perform, and the saved state.", 10)
    p(doc, "White-box design selects paths after reading the code. Each graph states its scope. Cyclomatic complexity V(G) = E - N + 2 gives the number of independent paths; one test is written per path, and each path adds an edge the earlier paths did not use. Source-derived paths are not an instrumented coverage measurement.", 10)
    p(doc, f"Implementation baseline: {BASE_REVISION} (main, 20 October 2026). UC1-04 deleteAccount is the control class chosen for black-box testing, as required by Lab 4 section 3.2.1; UC1-03 UpdateProfile is a second executed table. Lab 4 section 3.2.2 asks for basis paths of two methods; this report covers one method for each use case (five in total).", 9)
    p(doc, "Format follows the team's session testing report (docs/testing/session-testing).", 9)

    for group in cases:
        heading(doc, f"Black box testing {group['uc']}")
        doc.add_heading(group["title"], level=2)
        p(doc, f"Scenario: {group['scenario']}  Priority: High.", 9.5)
        p(doc, "Prerequisites: " + group["prerequisites"], 9.5)
        rows = [[r["id"], r["description"], r["action"], r["inputs"], r["expected"], r["actual"], r["result"]] for r in group["cases"]]
        table(doc, ["ID", "Description", "Action", "Inputs", "Expected output", "Actual output", "Result"],
              rows, [0.45, 0.95, 1.3, 1.3, 2.55, 3.15, 0.78], 8.2)
        p(doc, f"Run on {production['target']} with {production['browser']}. {production['accounts']} "
               "Observations: evidence/production/browser-run-20261020.json.", 8.5)
        p(doc, "Source: " + group["source"], 8.5)

    heading(doc, "Defects found in the browser run")
    table(doc, ["ID", "Case", "Defect", "Status"], [
        ["D1", "UC1-02 (found while logging in)", "If Log in is pressed before the page's JavaScript has loaded, the browser submits the form as a plain GET, so the email and password appear in the address bar, browser history and server logs. The login and registration forms set no method.", "Open. Fix: give both forms method=\"post\"."],
        ["D2", "L05", "/payouts redirects to /login without ?next=/payouts, so the user lands on Home after logging in.", "Fixed in PR #98."],
        ["D3", "M07", "The Help page does not link to Payouts.", "Fixed in PR #98."],
        ["D4", "M01", "Settings shows no email address and has no Email or Payouts link (Email is reached through Help).", "Open."],
    ], [0.45, 1.6, 6.0, 2.43], 9)
    p(doc, "Expected outputs corrected after the run: G02 now states that the name field holds at most 60 characters, because the 61st character cannot be typed. "
           "Observations that need no change: the duplicate-email check appears on step 2, when the account is created (R05); a deleted member stays in the member list as \"Deleted user\", as the anonymise-and-retain rule requires (X04).", 9)
    shots = HERE / "evidence/production/screenshots"
    for name, caption in [
        ("R05-existing-email-account-A", "R05: duplicate email refused, with a Log in instead link."),
        ("X01-owns-two-groups", "X01: owned groups block deletion; the delete button is not offered."),
        ("G05-old-link-while-off-member-C", "G05: a turned-off invitation link is refused."),
        ("X04-account-C-nothing-outstanding", "X04: after deletion the login page confirms the account was deleted."),
    ]:
        figure(doc, shots / f"{name}.jpg", max_height=2.4, description=caption)
        p(doc, caption, 9)

    executed_tables(doc, "Executed black-box tests: delete account", "UC1-04 deleteAccount", delete_obs, delete_input, delete_output,
                    "Baseline: active account with a wallet and nothing outstanding (all six obligations 0). Money is in integer cents. "
                    "Outputs list the outcome, the port steps performed, and the profile status afterwards. INVALID_INPUT means the obligation facts themselves were rejected.",
                    "evidence/delete-account-black-box-observations.json; tests/lab4/delete-account-black-box.test.ts.")
    executed_tables(doc, "Executed black-box tests: manage profile", "UC1-03 UpdateProfile", profile_obs, profile_input, profile_output,
                    "Baseline: active account; name \"Marcus Lim\"; sports [Badminton]; regions [West]. Supported sports: Badminton, Basketball, Football, Futsal, Tennis, Volleyball. "
                    "Supported regions: Central, East, North, North-East, West. Names are trimmed before the 1-60 character check.",
                    "evidence/update-profile-black-box-observations.json; tests/lab4/update-profile-black-box.test.ts.")

    for item in white:
        graph = graphs[item["graph"]]
        assert len(item["paths"]) == graph["counts"]["cyclomaticComplexity"]
        for case, path in zip(item["paths"], graph["paths"], strict=True):
            assert case[1] == " ".join(path["nodes"])
            observed = basis_obs["WB-" + case[0]]
            assert observed["predictedPath"] == case[1] and observed["expected"] == observed["actual"]
        heading(doc, f"White box testing {item['uc']}")
        doc.add_heading(item["title"], level=2)
        p(doc, "Source: " + item["source"], 9)
        p(doc, item["scope"], 9.5)
        figure(doc, HERE / "figures" / item["image"], max_height=2.3)
        p(doc, item["complexity"], 10, bold=True)
        heading(doc, f"Basis paths for {item['uc']}")
        p(doc, item["title"], 11, bold=True)
        table(doc, ["Path", "Basis set of nodes", "Input and expected output", "Execution evidence"], item["paths"],
              [0.55, 2.6, 4.4, 2.93], 9)
        for note in item["notes"]:
            p(doc, note, 9.3)

    heading(doc, "Execution evidence and reproduction")
    p(doc, f"All {totals['passed']} focused checks passed; {totals['todo']} are recorded as todo and none failed or were skipped. "
           f"Execution used Node {metadata['node']} and Vitest {metadata['vitest']} on 20 October 2026 ({metadata['runAt']} UTC). "
           "The report reuses the recorded outputs; no manual browser result is inferred from those checks.", 10)
    purpose = {
        "tests/lab4/delete-account-black-box.test.ts": "Lab 4 UC1-04 equivalence classes and boundary values",
        "tests/lab4/update-profile-black-box.test.ts": "Lab 4 UC1-03 equivalence classes and boundary values",
        "tests/lab4/account-basis-path.test.ts": "Lab 4 basis paths: deleteAccount and group join",
        "tests/lab4/auth-profile-basis-path.test.ts": "Lab 4 basis paths: registerUser, safeRedirectPath, prepareProfileUpdate",
        "tests/use-cases/UC1-01-register-user.test.ts": "Registration form, sign-up action and outcomes",
        "tests/use-cases/UC1-02-authenticate-user.test.ts": "Login, middleware protection, safe redirects",
        "tests/use-cases/UC1-03-manage-profile.test.ts": "Profile policy and saving",
        "tests/use-cases/UC1-04-delete-account.test.ts": "Deletion order, undo and blockers",
        "tests/use-cases/UC1-06-manage-group.test.ts": "Group create, join, members, links and archive",
        "tests/domain/accounts/deactivation-policy.test.ts": "Obligation policy",
        "tests/domain/groups/regular-group.test.ts": "Group aggregate rules",
        "tests/app/profile/actions.test.ts": "Profile server action",
        "tests/app/profile/settings-links.test.ts": "Settings and Help links",
    }
    rows = []
    for result in run["testResults"]:
        name = result["name"].replace("<repo>/", "")
        counts = [a["status"] for a in result["assertionResults"]]
        label = str(counts.count("passed")) + (f" (+{counts.count('todo')} todo)" if "todo" in counts else "")
        rows.append([name, label, purpose[name]])
    rows.append(["Total", f"{totals['passed']} (+{totals['todo']} todo)", f"{totals['files']} files; in-memory and mocked ports only"])
    table(doc, ["Executed suite", "Tests", "What the result establishes"], rows, [4.3, 1.2, 4.98], 9)
    p(doc, "Recorded todos: " + "; ".join(metadata["todo"]) + ". The first matches a known gap at this baseline: Settings calculates the reliability score from an empty history (fix proposed in PR #101).", 9)
    p(doc, "Automated tests use in-memory ports and mocks. The browser scenarios cover the deployed pages and Supabase Auth; PostgreSQL row locking, email delivery and the money-dependent blockers remain untested in the browser. No branch-coverage percentage is claimed.", 9)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=HERE / "account-testing-report.docx")
    create_report(parser.parse_args().output)
